"""
Import intelligent de produits : Excel, CSV, JSON, PDF, Word, texte, archive ZIP et images, en une seule fois.

Etapes (chaque requete reste courte, meme avec des dizaines de photos) :
  analyze -> lit les documents, detecte les colonnes, associe les photos aux produits et renvoie un apercu (aucune ecriture)
  rows    -> cree / met a jour les produits, categories et stocks
  images  -> rattache les photos par petits lots (redimensionnees) aux produits
"""

import csv
import difflib
import io
import ipaddress
import json
import os
import re
import shutil
import socket
import tempfile
import time
import unicodedata
import uuid
import zipfile
from decimal import Decimal, InvalidOperation
from urllib.parse import urlparse

import requests
from django.core.files.base import ContentFile
from django.db import transaction
from PIL import Image

from apps.stores.models import PointOfSale, Stock, StockMovement, StoreCategory
from apps.stores.services import change_stock

from .models import Category, Product

IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".jfif", ".bmp")
TABLE_EXTS = (".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".json", ".pdf", ".docx")
MAX_IMAGE_BYTES = 12 * 1024 * 1024
JOB_ROOT = os.path.join(tempfile.gettempdir(), "lbt_import")

# nom de colonne (sans accents, minuscules, sans ponctuation) -> champ
ALIASES = {
    "sku": ["sku", "reference", "ref", "code", "codeproduit", "codearticle", "codebarre", "codebarres", "ean", "id", "numero"],
    "name": ["name", "nom", "produit", "designation", "article", "libelle", "titre", "plat", "nomduproduit", "nomproduit", "description produit"],
    "category": ["category", "categorie", "famille", "rayon", "type", "groupe", "section"],
    "price": ["price", "prix", "prixunitaire", "prixvente", "prixdevente", "pu", "tarif", "montant", "prixttc", "prixfcfa", "prixxof"],
    "compare_at_price": ["compareatprice", "prixbarre", "ancienprix", "prixinitial"],
    "unit": ["unit", "unite", "conditionnement"],
    "description": ["description", "details", "detail", "commentaire", "note", "composition"],
    "is_active": ["isactive", "actif", "active", "visible", "statut", "disponible"],
    "image_url": ["imageurl", "image", "lienimage", "urlimage", "photo", "lienphoto", "urlphoto", "lien"],
    "image_file": ["fichierimage", "fichierphoto", "nomimage", "nomphoto", "imagefile", "photofichier"],
    "stock": ["stock", "quantite", "qte", "qty", "quantity", "stockdisponible", "inventaire"],
}
ALIAS_LOOKUP = {}
for field, names in ALIASES.items():
    for n in names:
        ALIAS_LOOKUP.setdefault(re.sub(r"[^a-z0-9]", "", n.lower()), field)


# ------------------------------------------------------------------ utilitaires
def norm(text):
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def compact(text):
    return norm(text).replace(" ", "")


def parse_price(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    s = re.sub(r"(?i)(fcfa|f\s*cfa|cfa|xof|francs?|\bf\b)", "", str(value)).replace(" ", " ").strip()
    s = re.sub(r"[^\d.,\s-]", "", s).replace(" ", "")
    if not s or not re.search(r"\d", s):
        return None
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", s):  # 1.500 ou 1,500 ou 1.500.000 : separateur de milliers
        s = re.sub(r"[.,]", "", s)
    elif "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def parse_bool(value, default=True):
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip().lower() in ("oui", "yes", "true", "1", "actif", "active", "vrai", "x", "disponible")


def parse_int(value):
    try:
        return int(Decimal(str(value).replace(" ", "").replace(",", ".")))
    except (InvalidOperation, ValueError):
        return None


def detect_header(rows):
    """Cherche la ligne d'en-tetes (au moins 'nom' + un autre champ connu) dans les 12 premieres lignes."""
    for i, row in enumerate(rows[:12]):
        mapping = {}
        for j, cell in enumerate(row):
            field = ALIAS_LOOKUP.get(re.sub(r"[^a-z0-9]", "", norm(cell).replace(" ", "")))
            if field and field not in mapping.values():
                mapping[j] = field
        if "name" in mapping.values() and len(mapping) >= 2:
            return i, mapping
    return None, {}


def rows_to_products(rows, source, warnings, xlsx_images=None):
    """Transforme des lignes brutes (listes de cellules) en dictionnaires produits normalises."""
    rows = [list(r) for r in rows if r and any(str(c).strip() for c in r if c is not None)]
    if not rows:
        return []
    hi, mapping = detect_header(rows)
    if hi is None:
        # pas d'en-tete : tableau a 2 colonnes "nom, prix" ou "nom, categorie, prix"
        width = max(len(r) for r in rows)
        if width == 2:
            mapping, hi = {0: "name", 1: "price"}, -1
        elif width == 3:
            mapping, hi = {0: "name", 1: "category", 2: "price"}, -1
        else:
            warnings.append(f"{source} : colonnes non reconnues (il faut au moins une colonne « nom » et une colonne « prix »).")
            return []
    out = []
    for offset, row in enumerate(rows[hi + 1 :], start=hi + 2):
        item = {"source": source, "line": offset}
        for j, field in mapping.items():
            if j < len(row) and row[j] is not None and str(row[j]).strip() != "":
                item[field] = row[j].strip() if isinstance(row[j], str) else row[j]
        if not str(item.get("name") or "").strip():
            continue
        item["name"] = re.sub(r"\s+", " ", str(item["name"])).strip()
        if xlsx_images and (offset - 1) in xlsx_images:
            item["embedded"] = xlsx_images[offset - 1]
        out.append(item)
    return out


LINE_RE = re.compile(r"^\s*(?:\d+\s*[.)-]\s+)?(?P<name>[^\d\n][^\n]*?)\s*[-:.…–_ ]{1,}\s*(?P<price>\d[\d\s.,]*)\s*(?:fcfa|cfa|f|xof)?\s*$", re.I)


def text_to_products(text, source, warnings):
    """Menus / listes de prix en texte libre : « Yassa poulet .... 1 500 F », les titres en majuscules deviennent des categories."""
    out, category = [], None
    for n, line in enumerate(text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        m = LINE_RE.match(line)
        if m and re.search(r"[A-Za-zÀ-ſ]{2}", m.group("name")) and parse_price(m.group("price")) is not None:
            item = {"source": source, "line": n, "name": re.sub(r"\s+", " ", m.group("name")).strip(" .-:"), "price": m.group("price")}
            if category:
                item["category"] = category
            out.append(item)
        elif len(line) <= 40 and line.upper() == line and re.search(r"[A-Z]{3}", line):
            category = line.title()
    if not out:
        warnings.append(f"{source} : aucune ligne « produit + prix » reconnue.")
    return out


# ------------------------------------------------------------------ lecture des documents
def read_xlsx(data, name, warnings):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    products, images = [], []
    for ws in wb.worksheets:
        if ws.title.lower() in ("instructions", "exemple", "example"):
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # images inserees dans les cellules : associees a la ligne ou elles se trouvent
        by_row = {}
        for img in getattr(ws, "_images", []):
            try:
                row0 = img.anchor._from.row  # 0-based
                blob = img._data()
                by_row[row0] = len(images)
                images.append((f"{ws.title}-ligne{row0 + 1}.png", blob))
            except Exception:  # noqa: BLE001
                continue
        products += rows_to_products(rows, f"{name} ({ws.title})", warnings, xlsx_images=by_row)
    return products, images


def read_csv(data, name, warnings):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = list(csv.reader(io.StringIO(text), dialect))
    products = rows_to_products(rows, name, warnings)
    return products or text_to_products(text, name, warnings)


def read_json(data, name, warnings):
    obj = json.loads(data.decode("utf-8-sig"))
    if isinstance(obj, dict):
        obj = next((v for v in obj.values() if isinstance(v, list)), [])
    out = []
    for n, entry in enumerate(obj, start=1):
        if not isinstance(entry, dict):
            continue
        item = {"source": name, "line": n}
        for k, v in entry.items():
            field = ALIAS_LOOKUP.get(compact(k))
            if field and v not in (None, ""):
                item[field] = v
        if str(item.get("name") or "").strip():
            out.append(item)
    return out


def read_pdf(data, name, warnings):
    try:
        import pdfplumber
    except ImportError:
        warnings.append(f"{name} : lecture des PDF indisponible sur ce serveur.")
        return []
    products, text_all = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages[:60]:
            for table in page.extract_tables() or []:
                products += rows_to_products(table, f"{name} (p.{page.page_number})", warnings)
            text_all.append(page.extract_text() or "")
    return products or text_to_products("\n".join(text_all), name, warnings)


def read_docx(data, name, warnings):
    try:
        import docx
    except ImportError:
        warnings.append(f"{name} : lecture des documents Word indisponible sur ce serveur.")
        return []
    document = docx.Document(io.BytesIO(data))
    products = []
    for t in document.tables:
        products += rows_to_products([[c.text for c in r.cells] for r in t.rows], name, warnings)
    return products or text_to_products("\n".join(p.text for p in document.paragraphs), name, warnings)


def read_document(name, data, warnings):
    ext = os.path.splitext(name.lower())[1]
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(data, name, warnings)
    if ext in (".csv", ".tsv"):
        return read_csv(data, name, warnings), []
    if ext == ".txt":
        text = data.decode("utf-8-sig", errors="replace")
        return (read_csv(data, name, warnings) if text.count(";") + text.count("\t") > 3 else text_to_products(text, name, warnings)), []
    if ext == ".json":
        return read_json(data, name, warnings), []
    if ext == ".pdf":
        return read_pdf(data, name, warnings), []
    if ext == ".docx":
        return read_docx(data, name, warnings), []
    if ext in (".xls", ".ods", ".doc"):
        warnings.append(f"{name} : ancien format non pris en charge. Enregistrez-le en .xlsx (Excel) ou .docx puis reessayez.")
    return [], []


# ------------------------------------------------------------------ etape 1 : analyse
def _cleanup_old_jobs():
    if not os.path.isdir(JOB_ROOT):
        return
    for d in os.listdir(JOB_ROOT):
        path = os.path.join(JOB_ROOT, d)
        if time.time() - os.path.getmtime(path) > 6 * 3600:
            shutil.rmtree(path, ignore_errors=True)


def _job_dir(job_id):
    if not re.fullmatch(r"[0-9a-f]{32}", job_id or ""):
        raise ValueError("Import introuvable.")
    return os.path.join(JOB_ROOT, job_id)


def _load(job_id):
    path = os.path.join(_job_dir(job_id), "plan.json")
    if not os.path.exists(path):
        raise ValueError("Import expire : recommencez l'analyse.")
    return json.load(open(path, encoding="utf-8"))


def _save(job_id, plan):
    json.dump(plan, open(os.path.join(_job_dir(job_id), "plan.json"), "w", encoding="utf-8"), ensure_ascii=False)


def match_image(item, pool, used):
    """Renvoie l'index de l'image du lot qui correspond au produit : fichier indique, SKU, nom exact, puis nom approchant."""
    if "embedded" in item and item["embedded"] not in used:
        return item["embedded"]
    stems = [(i, os.path.splitext(img["filename"])[0], img["filename"]) for i, img in enumerate(pool) if i not in used and not img.get("embedded")]
    wanted = str(item.get("image_file") or "").strip().lower()
    if wanted:
        for i, _stem, fn in stems:
            if fn.lower() == wanted or os.path.splitext(fn.lower())[0] == os.path.splitext(wanted)[0]:
                return i
    sku = compact(item.get("sku", ""))
    if sku:
        for i, stem, _ in stems:
            if compact(stem) == sku:
                return i
    cname = compact(item["name"])
    for i, stem, _ in stems:
        if compact(stem) == cname:
            return i
    best, best_i = 0.0, None
    for i, stem, _ in stems:
        cs = compact(stem)
        if len(cs) >= 4 and (cs in cname or cname in cs) and min(len(cs), len(cname)) >= 4:
            ratio = 0.9
        else:
            ratio = difflib.SequenceMatcher(None, cs, cname).ratio()
        if ratio > best:
            best, best_i = ratio, i
    return best_i if best >= 0.84 else None


def analyze(files, store=None, create_from_images=False):
    """files : liste de (nom, octets). Renvoie l'apercu ; les fichiers utiles sont gardes sur disque pour la suite."""
    _cleanup_old_jobs()
    job_id = uuid.uuid4().hex
    os.makedirs(_job_dir(job_id), exist_ok=True)
    warnings, products, pool, kinds = [], [], [], []

    def add_image(name, data, embedded=False):
        try:
            with Image.open(io.BytesIO(data)) as im:
                im.verify()
        except Exception:  # noqa: BLE001
            warnings.append(f"{name} : image illisible, ignoree.")
            return
        if len(data) > MAX_IMAGE_BYTES:
            warnings.append(f"{name} : image trop lourde (> 12 Mo), ignoree.")
            return
        idx = len(pool)
        fn = f"img_{idx}{os.path.splitext(name.lower())[1] or '.jpg'}"
        open(os.path.join(_job_dir(job_id), fn), "wb").write(data)
        pool.append({"filename": os.path.basename(name), "file": fn, "embedded": embedded})
        return idx

    def handle(name, data, depth=0):
        ext = os.path.splitext(name.lower())[1]
        base = os.path.basename(name)
        if base.startswith(("._", "~$")) or "__MACOSX" in name:
            return
        if ext == ".zip" and depth < 2:
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    for info in z.infolist()[:2000]:
                        if not info.is_dir() and info.file_size <= 60 * 1024 * 1024:
                            handle(info.filename, z.read(info), depth + 1)
                kinds.append({"name": name, "kind": "archive ZIP"})
            except zipfile.BadZipFile:
                warnings.append(f"{name} : archive ZIP illisible.")
        elif ext in IMAGE_EXTS:
            add_image(name, data)
        elif ext in TABLE_EXTS or ext in (".xls", ".ods", ".doc"):
            try:
                items, embedded = read_document(name, data, warnings)
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"{name} : lecture impossible ({str(exc)[:120]}).")
                return
            emb_index = {}
            for k, (iname, blob) in enumerate(embedded):
                idx = add_image(iname, blob, embedded=True)
                if idx is not None:
                    emb_index[k] = idx
            for it in items:
                if "embedded" in it:
                    it["embedded"] = emb_index.get(it["embedded"])
                    if it["embedded"] is None:
                        del it["embedded"]
            products.extend(items)
            kinds.append({"name": name, "kind": f"{len(items)} produit(s)"})
        else:
            warnings.append(f"{base} : type de fichier non pris en charge.")

    for name, data in files:
        handle(name, data)

    # associer chaque ligne a un produit existant (SKU puis nom) et a une photo
    by_sku = {p.sku.lower(): p for p in Product.objects.all()}
    # avec un point de vente choisi, un nom identique ne rapproche que les produits DE CE point de vente (pas ceux des autres)
    scope = Product.objects.filter(stocks__point_of_sale=store).distinct() if store else Product.objects.all()
    by_name = {norm(p.name): p for p in scope}
    used, seen, rows = set(), set(), []
    for it in products:
        sku = str(it.get("sku") or "").strip()
        existing = by_sku.get(sku.lower()) if sku else None
        existing = existing or by_name.get(norm(it["name"]))
        key = (existing.pk if existing else norm(it["name"]))
        if key in seen:
            warnings.append(f"« {it['name']} » apparait plusieurs fois : seule la premiere ligne est importee.")
            continue
        seen.add(key)
        price = parse_price(it.get("price"))
        row = {
            "line": it["line"],
            "source": it["source"],
            "sku": sku,
            "name": it["name"],
            "category": str(it.get("category") or "").strip(),
            "price": str(price) if price is not None else None,
            "compare_at_price": str(parse_price(it.get("compare_at_price"))) if parse_price(it.get("compare_at_price")) is not None else None,
            "unit": str(it.get("unit") or "").strip(),
            "description": str(it.get("description") or "").strip(),
            "is_active": parse_bool(it.get("is_active"), default=(price is not None and price > 0)),
            "stock": parse_int(it.get("stock")) if it.get("stock") not in (None, "") else None,
            "image_url": str(it.get("image_url") or "").strip() if str(it.get("image_url") or "").lower().startswith("http") else "",
            "action": "update" if existing else "create",
            "product_id": existing.pk if existing else None,
        }
        if not row["image_url"]:
            if it.get("image_url") and not row.get("image_file") and "image_file" not in it:
                it["image_file"] = it.get("image_url")  # colonne « photo » contenant un nom de fichier
            m = match_image(it, pool, used)
            if m is not None:
                used.add(m)
                row["image_index"] = m
        rows.append(row)
        if price is None:
            row["warning"] = "Prix manquant : le produit sera cree masque (prix 0)."

    unmatched = [pool[i]["filename"] for i in range(len(pool)) if i not in used]
    if create_from_images:
        for i in range(len(pool)):
            if i in used:
                continue
            title = re.sub(r"[-_]+", " ", os.path.splitext(pool[i]["filename"])[0]).strip().capitalize()
            existing = by_name.get(norm(title))
            rows.append(
                {
                    "line": 0, "source": pool[i]["filename"], "sku": "", "name": title, "category": "", "price": None, "compare_at_price": None,
                    "unit": "", "description": "", "is_active": False, "stock": None, "image_url": "", "image_index": i,
                    "action": "update" if existing else "create", "product_id": existing.pk if existing else None,
                    "warning": "Cree a partir de la photo : prix a renseigner (masque en attendant).",
                }
            )
            used.add(i)
        unmatched = []

    plan = {"store_id": store.pk if store else None, "rows": rows, "images": pool, "stage": "analyzed", "pending": [], "stats": {}}
    _save(job_id, plan)
    return {
        "job_id": job_id,
        "documents": kinds,
        "warnings": warnings[:30],
        "total_rows": len(rows),
        "create": sum(1 for r in rows if r["action"] == "create"),
        "update": sum(1 for r in rows if r["action"] == "update"),
        "images_total": len(pool),
        "images_matched": len(used),
        "images_unmatched": unmatched[:60],
        "image_names": [img["filename"] for img in pool],
        "rows": rows[:300],
    }


# ------------------------------------------------------------------ etape 2 : produits
def get_category(name, cache):
    """Categorie existante (sans tenir compte des majuscules ni des accents) ou nouvelle categorie."""
    key = norm(name)
    if key not in cache:
        found = next((c for c in Category.objects.all() if norm(c.name) == key), None)
        cache[key] = found or Category.objects.create(name=name.strip())
    return cache[key]


def _next_sku(prefix, taken):
    n = 1
    while True:
        sku = f"{prefix}-{n:04d}"
        if sku.lower() not in taken:
            taken.add(sku.lower())
            return sku
        n += 1


def run_rows(job_id):
    plan = _load(job_id)
    store = PointOfSale.objects.filter(pk=plan["store_id"]).first() if plan["store_id"] else None
    prefix = (store.slug or "imp").upper()[:6] if store else "IMP"
    taken = {s.lower() for s in Product.objects.values_list("sku", flat=True)}
    created = updated = 0
    errors, pending, cats = [], [], {}
    for row in plan["rows"]:
        try:
            with transaction.atomic():
                product = Product.objects.filter(pk=row["product_id"]).first() if row.get("product_id") else None
                price = Decimal(row["price"]) if row.get("price") is not None else None
                category = get_category(row["category"], cats) if row.get("category") else None
                if product is None:
                    product = Product(sku=row["sku"] or _next_sku(prefix, taken), name=row["name"], price=price or Decimal(0))
                    product.is_active = row["is_active"] and (price or 0) > 0
                    created += 1
                else:
                    product.name = row["name"] or product.name
                    if price is not None:
                        product.price = price
                    updated += 1
                    if "is_active" in row and row.get("price") is not None:
                        product.is_active = row["is_active"]
                if category:
                    product.category = category
                if row.get("compare_at_price"):
                    product.compare_at_price = Decimal(row["compare_at_price"])
                if row.get("unit"):
                    product.unit = row["unit"][:32]
                if row.get("description"):
                    product.description = row["description"]
                product.save()
                if store:
                    Stock.objects.get_or_create(product=product, point_of_sale=store)
                    if product.category_id:
                        StoreCategory.objects.get_or_create(
                            point_of_sale=store,
                            category_id=product.category_id,
                            defaults={"order": StoreCategory.objects.filter(point_of_sale=store).count()},
                        )
                    if row.get("stock") is not None:
                        change_stock(product, store, set_to=max(row["stock"], 0), reason=StockMovement.Reason.IMPORT)
                if row.get("image_index") is not None:
                    pending.append({"product_id": product.pk, "kind": "file", "file": plan["images"][row["image_index"]]["file"]})
                elif row.get("image_url"):
                    pending.append({"product_id": product.pk, "kind": "url", "url": row["image_url"]})
        except Exception as exc:  # noqa: BLE001
            errors.append(f"Ligne {row.get('line') or row['name']} ({row['name']}) : {str(exc)[:160]}")
    plan.update({"stage": "rows_done", "pending": pending, "stats": {"created": created, "updated": updated, "images_attached": 0, "errors": errors[:50]}})
    _save(job_id, plan)
    return {"created": created, "updated": updated, "errors": errors[:50], "images_pending": len(pending)}


# ------------------------------------------------------------------ etape 3 : photos
def _safe_url(url):
    p = urlparse(url)
    if p.scheme not in ("http", "https") or not p.hostname:
        raise ValueError("adresse invalide")
    for info in socket.getaddrinfo(p.hostname, None):
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            raise ValueError("adresse non autorisee")


def _prepare_jpeg(data):
    with Image.open(io.BytesIO(data)) as im:
        im = im.convert("RGB")
        im.thumbnail((1600, 1600))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=88, optimize=True, progressive=True)
        return buf.getvalue()


def run_images(job_id, batch=4, budget=20.0):
    plan = _load(job_id)
    started, attached, errors = time.time(), 0, []
    pending = plan["pending"]
    while pending and attached + len(errors) < batch and time.time() - started < budget:
        item = pending.pop(0)
        try:
            product = Product.objects.get(pk=item["product_id"])
            if item["kind"] == "file":
                data = open(os.path.join(_job_dir(job_id), item["file"]), "rb").read()
            else:
                _safe_url(item["url"])
                resp = requests.get(item["url"], timeout=12, stream=True)
                resp.raise_for_status()
                data = resp.raw.read(MAX_IMAGE_BYTES + 1, decode_content=True)
                if len(data) > MAX_IMAGE_BYTES:
                    raise ValueError("image trop lourde")
            product.image.save(f"{product.sku}.jpg", ContentFile(_prepare_jpeg(data)), save=True)
            attached += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{item.get('url') or item.get('file')} : {str(exc)[:120]}")
    stats = plan["stats"]
    stats["images_attached"] = stats.get("images_attached", 0) + attached
    stats["errors"] = (stats.get("errors", []) + errors)[:50]
    plan["pending"] = pending
    if not pending:
        plan["stage"] = "done"
    _save(job_id, plan)
    if not pending:
        shutil.rmtree(_job_dir(job_id), ignore_errors=True)
    return {"attached": stats["images_attached"], "remaining": len(pending), "errors": stats["errors"], "done": not pending}
