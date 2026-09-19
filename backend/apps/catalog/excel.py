import io
from urllib.parse import urlparse

import openpyxl
import requests
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from django.core.files.base import ContentFile
from openpyxl.utils import get_column_letter

from apps.stores.models import PointOfSale, Stock, StockMovement, StoreCategory
from apps.stores.services import change_stock

from .models import Category, Product

BASE_HEADERS = [
    "sku",
    "name",
    "category",
    "price",
    "compare_at_price",
    "unit",
    "description",
    "is_active",
    "image_url",
]

STOCK_COLUMN_PREFIX = "stock:"


def export_products_to_excel(queryset=None, store=None):
    """
    Retourne un fichier .xlsx (bytes) listant les produits. Sans `store` : une colonne
    "stock:<Point de vente>" par point de vente actif. Avec `store` : uniquement les produits
    de ce point de vente, avec une seule colonne "stock" (le fichier se reimporte tel quel).
    """
    queryset = queryset if queryset is not None else Product.objects.select_related("category").prefetch_related(
        "stocks__point_of_sale"
    )
    if store:
        queryset = queryset.filter(stocks__point_of_sale=store).distinct()
        stores = [store]
        headers = BASE_HEADERS + ["stock"]
    else:
        stores = list(PointOfSale.objects.filter(is_active=True).order_by("name"))
        headers = BASE_HEADERS + [f"{STOCK_COLUMN_PREFIX}{s.name}" for s in stores]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produits"
    ws.append(headers)

    for product in queryset:
        stock_by_store = {s.point_of_sale_id: s.quantity for s in product.stocks.all()}
        row = [
            product.sku,
            product.name,
            product.category.name if product.category else "",
            float(product.price),
            float(product.compare_at_price) if product.compare_at_price else None,
            product.unit,
            product.description,
            "oui" if product.is_active else "non",
            product.image.url if product.image else "",
        ]
        row += [stock_by_store.get(store.id, 0) for store in stores]
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 4)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


HEADER_HELP = {
    "sku": "OBLIGATOIRE. Code unique du produit (ex. RIZ-25KG). Sert a retrouver le produit : meme SKU = mise a jour.",
    "name": "OBLIGATOIRE. Nom affiche aux clients.",
    "category": "Categorie (creee automatiquement si elle n'existe pas).",
    "price": "OBLIGATOIRE. Prix de vente en FCFA, nombre sans espace ni symbole.",
    "compare_at_price": "Ancien prix barre (optionnel). Doit etre superieur au prix.",
    "unit": "Unite de vente : piece, kg, litre, sac...",
    "description": "Description courte (optionnel).",
    "is_active": "oui = visible sur le site, non = masque.",
    "image_url": "Lien https vers la photo du produit (optionnel).",
}

EXAMPLE_ROWS = [
    ["RIZ-25KG", "Riz brise 25 kg", "Epicerie", 15500, 17000, "sac", "Riz brise de qualite superieure", "oui", "", 40, 10],
    ["HUILE-5L", "Huile vegetale 5 L", "Epicerie", 6500, None, "bidon", "", "oui", "", 25, 0],
    ["MAIS-50KG", "Mais concasse 50 kg", "Aliment volaille", 14000, None, "sac", "Aliment pour volaille", "oui", "", 0, 60],
]


def build_import_template(store=None):
    """
    Modele Excel d'import produits : feuille Produits (vide, a remplir) + Exemple + Instructions.
    Avec `store` : modele propre a un point de vente (une seule colonne "stock").
    """
    if store:
        stores = [store]
        headers = BASE_HEADERS + ["stock"]
    else:
        stores = list(PointOfSale.objects.filter(is_active=True).order_by("name"))
        headers = BASE_HEADERS + [f"{STOCK_COLUMN_PREFIX}{s.name}" for s in stores]
    red = PatternFill("solid", fgColor="8B1A1A")
    gold = PatternFill("solid", fgColor="D4A017")
    required = {"sku", "name", "price"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produits"
    ws.append(headers)
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = red if h in required else PatternFill("solid", fgColor="5B5B5B")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        help_text = HEADER_HELP.get(h) or (
            f"Quantite en stock pour ce point de vente ({h[len(STOCK_COLUMN_PREFIX):]})."
            if h.startswith(STOCK_COLUMN_PREFIX)
            else f"Quantite en stock a {store.name if store else 'ce point de vente'}."
        )
        cell.comment = Comment(help_text, "La Belle Teranga")
        ws.column_dimensions[get_column_letter(idx)].width = max(16, len(h) + 4)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    bool_dv = DataValidation(type="list", formula1='"oui,non"', allow_blank=True)
    ws.add_data_validation(bool_dv)
    bool_dv.add(f"H2:H2000")
    price_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    price_dv.error = "Le prix doit etre un nombre positif."
    ws.add_data_validation(price_dv)
    price_dv.add("D2:E2000")
    if stores:
        qty_dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        qty_dv.error = "La quantite doit etre un entier positif."
        ws.add_data_validation(qty_dv)
        qty_dv.add(f"{get_column_letter(len(BASE_HEADERS) + 1)}2:{get_column_letter(len(headers))}2000")

    example = wb.create_sheet("Exemple")
    example.append(headers)
    for idx in range(1, len(headers) + 1):
        c = example.cell(row=1, column=idx)
        c.font = Font(bold=True)
        c.fill = gold
        example.column_dimensions[get_column_letter(idx)].width = max(16, len(headers[idx - 1]) + 4)
    for row in EXAMPLE_ROWS:
        stock_values = (row[len(BASE_HEADERS):] + [0] * len(stores))[: len(stores)]
        example.append(row[: len(BASE_HEADERS)] + stock_values)

    info = wb.create_sheet("Instructions")
    lines = [
        ("MODELE D'IMPORT DES PRODUITS - La Belle Teranga", True),
        ("", False),
        ("1. Remplissez la feuille 'Produits' : une ligne par produit (ne changez pas l'ordre ni le nom des colonnes).", False),
        ("2. Les colonnes en rouge sont obligatoires : sku, name, price. Les colonnes grises sont facultatives.", False),
        ("3. Le SKU est le code unique du produit : importer un SKU existant met le produit a jour au lieu de le dupliquer.", False),
        *(
            [
                (f"4. Ce fichier est propre a : {store.name}. La colonne 'stock' est la quantite disponible dans CE point de vente.", False),
                ("   Tous les produits de la feuille sont rattaches a ce point de vente (et seulement a lui).", False),
            ]
            if store
            else [
                ("4. Une colonne 'stock:<Point de vente>' existe pour chaque point de vente actif : saisissez la quantite disponible.", False),
                ("   Le nom apres 'stock:' doit correspondre exactement au nom du point de vente dans l'application.", False),
            ]
        ),
        ("5. Une cellule de stock laissee vide ne modifie pas le stock ; 0 met le stock a zero.", False),
        ("6. Les photos : indiquez un lien https dans image_url ; l'image est telechargee a l'import.", False),
        ("7. Importez le fichier depuis Admin > Produits, apres avoir choisi le point de vente, puis Importer Excel.", False),
        ("8. La feuille 'Exemple' montre 3 produits remplis : ne l'importez pas, seule la premiere feuille est lue.", False),
        ("", False),
        ("Points de vente actifs : " + (", ".join(s.name for s in stores) if stores else "aucun"), False),
    ]
    for text, bold in lines:
        info.append([text])
        info.cell(row=info.max_row, column=1).font = Font(bold=bold, size=13 if bold else 11)
    info.column_dimensions["A"].width = 120

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _attach_image_from_url(product, url):
    """Telecharge une image depuis une URL et la rattache au produit."""
    response = requests.get(url, timeout=15)
    response.raise_for_status()

    filename = urlparse(url).path.rsplit("/", 1)[-1] or f"{product.sku}.jpg"
    if "." not in filename:
        filename += ".jpg"

    product.image.save(filename, ContentFile(response.content), save=True)


def import_products_from_excel(file_obj, store=None):
    """
    Lit un fichier .xlsx et cree/met a jour les produits (upsert par sku).
    Les colonnes "stock:<Point de vente>" mettent a jour le stock de ce
    produit pour le point de vente correspondant (doit deja exister -
    utiliser /admin/stores pour creer un point de vente au prealable).
    Avec `store` (import propre a un point de vente) : chaque produit du fichier est rattache a ce
    point de vente (ligne de stock creee meme sans quantite) et la colonne "stock" (ou
    "stock:<nom du point de vente>") fixe sa quantite ; les colonnes des autres magasins sont ignorees.
    Retourne un resume: {created, updated, errors: [...]}
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "errors": ["Fichier vide."]}

    header = [str(h).strip() if h else "" for h in rows[0]]
    header_lower = [h.lower() for h in header]
    required = {"sku", "name", "price"}
    if not required.issubset(set(header_lower)):
        return {
            "created": 0,
            "updated": 0,
            "errors": [f"Colonnes obligatoires manquantes: {required - set(header_lower)}"],
        }

    col_index = {name: idx for idx, name in enumerate(header_lower)}
    category_cache = {}
    stock_columns = [
        (idx, h[len(STOCK_COLUMN_PREFIX):].strip())
        for idx, h in enumerate(header)
        if h.lower().startswith(STOCK_COLUMN_PREFIX)
    ]
    stores_by_name = {s.name.lower(): s for s in PointOfSale.objects.all()}
    store_stock_idx = None
    if store:
        store_stock_idx = col_index.get("stock")
        if store_stock_idx is None:
            store_stock_idx = next((i for i, n in stock_columns if n.lower() == store.name.lower()), None)
        stock_columns = []

    created, updated, errors = 0, 0, []

    for row_number, row in enumerate(rows[1:], start=2):
        if row is None or all(cell is None for cell in row):
            continue
        try:
            sku = str(row[col_index["sku"]]).strip()
            name = str(row[col_index["name"]]).strip()
            price = row[col_index["price"]]
            if not sku or not name or price is None:
                errors.append(f"Ligne {row_number}: sku, name et price sont obligatoires.")
                continue

            category_name = None
            if "category" in col_index and row[col_index["category"]]:
                category_name = str(row[col_index["category"]]).strip()

            category = None
            if category_name:
                from .smart_import import get_category  # insensible aux accents et aux majuscules : evite les doublons de categories

                category = get_category(category_name, category_cache)

            defaults = {
                "name": name,
                "price": price,
                "category": category,
            }
            if "compare_at_price" in col_index and row[col_index["compare_at_price"]] not in (None, ""):
                defaults["compare_at_price"] = row[col_index["compare_at_price"]]
            if "unit" in col_index and row[col_index["unit"]]:
                defaults["unit"] = str(row[col_index["unit"]]).strip()
            if "description" in col_index and row[col_index["description"]]:
                defaults["description"] = str(row[col_index["description"]]).strip()
            if "is_active" in col_index and row[col_index["is_active"]] is not None:
                val = str(row[col_index["is_active"]]).strip().lower()
                defaults["is_active"] = val in ("oui", "yes", "true", "1")

            obj, was_created = Product.objects.update_or_create(sku=sku, defaults=defaults)
            if was_created:
                created += 1
            else:
                updated += 1

            if "image_url" in col_index and row[col_index["image_url"]]:
                image_url = str(row[col_index["image_url"]]).strip()
                try:
                    _attach_image_from_url(obj, image_url)
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"Ligne {row_number}: image non recuperee ({image_url}) - {exc}")

            if store:
                Stock.objects.get_or_create(product=obj, point_of_sale=store)
                if obj.category_id:
                    StoreCategory.objects.get_or_create(
                        point_of_sale=store,
                        category_id=obj.category_id,
                        defaults={"order": StoreCategory.objects.filter(point_of_sale=store).count()},
                    )
                if store_stock_idx is not None and row[store_stock_idx] not in (None, ""):
                    change_stock(obj, store, set_to=int(row[store_stock_idx]), reason=StockMovement.Reason.IMPORT)

            for idx, store_name in stock_columns:
                if row[idx] is None or row[idx] == "":
                    continue
                store = stores_by_name.get(store_name.lower())
                if not store:
                    errors.append(
                        f"Ligne {row_number}: point de vente '{store_name}' introuvable, stock ignore "
                        f"(creez-le d'abord dans Points de vente)."
                    )
                    continue
                change_stock(obj, store, set_to=int(row[idx]), reason=StockMovement.Reason.IMPORT)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"Ligne {row_number}: {exc}")

    return {"created": created, "updated": updated, "errors": errors}
