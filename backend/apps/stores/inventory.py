import io

import openpyxl
from django.db import transaction
from django.utils import timezone
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

from apps.catalog.models import Product

from .models import InventoryCount, InventoryCountLine, Stock, StockMovement
from .services import change_stock


@transaction.atomic
def create_inventory(store, date, notes="", category=None, user=None):
    """Cree un inventaire et fige le stock theorique de tous les produits actifs."""
    from apps.documents.models import next_number  # import tardif : evite un cycle d'apps

    count = InventoryCount.objects.create(
        number=next_number("inventory", "INV"), point_of_sale=store, date=date, notes=notes, created_by=user
    )
    products = Product.objects.filter(is_active=True)
    if category:
        products = products.filter(category_id=category)
    stock = {s.product_id: s.quantity for s in Stock.objects.filter(point_of_sale=store)}
    InventoryCountLine.objects.bulk_create(
        [InventoryCountLine(count=count, product=p, theoretical_quantity=stock.get(p.id, 0)) for p in products]
    )
    return count


def ensure_draft(count):
    if count.status != InventoryCount.Status.DRAFT:
        raise ValidationError("Cet inventaire n'est plus modifiable.")


@transaction.atomic
def set_counts(count, counts):
    """counts = [{"product": id, "counted_quantity": int | None}, ...]"""
    ensure_draft(count)
    lines = {l.product_id: l for l in count.lines.select_related("product")}
    for entry in counts:
        line = lines.get(entry.get("product"))
        if not line:
            continue
        value = entry.get("counted_quantity")
        if value in (None, ""):
            line.counted_quantity = None
        else:
            try:
                value = int(value)
            except (TypeError, ValueError):
                raise ValidationError(f"Quantite invalide pour {line.product.name}.")
            if value < 0:
                raise ValidationError(f"Quantite negative pour {line.product.name}.")
            line.counted_quantity = value
        line.save(update_fields=["counted_quantity"])


@transaction.atomic
def validate_inventory(count, user=None):
    """Applique les quantites comptees au stock reel et journalise chaque ecart."""
    count = InventoryCount.objects.select_for_update().get(pk=count.pk)
    ensure_draft(count)
    counted = list(count.lines.select_related("product").exclude(counted_quantity__isnull=True))
    if not counted:
        raise ValidationError("Aucune quantite comptee : saisissez au moins un comptage.")
    moved = 0
    for line in counted:
        current = Stock.objects.filter(product=line.product, point_of_sale=count.point_of_sale).first()
        if (current.quantity if current else 0) != line.theoretical_quantity:
            moved += 1
        change_stock(
            line.product,
            count.point_of_sale,
            set_to=line.counted_quantity,
            reason=StockMovement.Reason.INVENTORY,
            reference=count.number,
            user=user,
        )
    count.status = InventoryCount.Status.VALIDATED
    count.validated_at = timezone.now()
    count.save(update_fields=["status", "validated_at"])
    return count, moved


COUNT_HEADERS = ["sku", "produit", "categorie", "stock_theorique", "quantite_comptee", "ecart"]


def export_counting_sheet(count):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comptage"
    ws.append(COUNT_HEADERS)
    fill = PatternFill("solid", fgColor="8B1A1A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for i, line in enumerate(count.lines.select_related("product__category"), start=2):
        ws.append(
            [
                line.product.sku,
                line.product.name,
                line.product.category.name if line.product.category else "",
                line.theoretical_quantity,
                line.counted_quantity,
                f'=IF(E{i}="","",E{i}-D{i})',
            ]
        )
    ws.freeze_panes = "A2"
    for idx, width in enumerate([16, 40, 20, 16, 18, 10], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    info = wb.create_sheet("Informations")
    for row in [
        ("Inventaire", count.number),
        ("Point de vente", count.point_of_sale.name),
        ("Date", str(count.date)),
        (
            "Consigne",
            "Remplissez la colonne quantite_comptee puis importez le fichier dans l'application. "
            "Laissez vide un produit non compte.",
        ),
    ]:
        info.append(row)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 90
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@transaction.atomic
def import_counts(count, file_obj):
    ensure_draft(count)
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception:  # noqa: BLE001
        raise ValidationError("Fichier Excel illisible (.xlsx attendu).")
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        raise ValidationError("Fichier vide.")
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    if "sku" not in header or "quantite_comptee" not in header:
        raise ValidationError("Colonnes sku et quantite_comptee obligatoires (utilisez le fichier exporte).")
    i_sku, i_qty = header.index("sku"), header.index("quantite_comptee")
    lines = {l.product.sku: l for l in count.lines.select_related("product")}
    updated, errors = 0, []
    for n, row in enumerate(rows[1:], start=2):
        if not row or row[i_sku] in (None, ""):
            continue
        line = lines.get(str(row[i_sku]).strip())
        if not line:
            errors.append(f"Ligne {n}: SKU {row[i_sku]} absent de cet inventaire.")
            continue
        value = row[i_qty]
        if value in (None, ""):
            continue
        try:
            value = int(float(value))
            if value < 0:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"Ligne {n}: quantite invalide ({value}).")
            continue
        line.counted_quantity = value
        line.save(update_fields=["counted_quantity"])
        updated += 1
    return {"updated": updated, "errors": errors}
