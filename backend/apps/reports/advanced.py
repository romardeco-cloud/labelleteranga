import io
from collections import defaultdict
from datetime import date as date_cls
from datetime import timedelta
from decimal import Decimal

import openpyxl
from django.db.models import Count, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.permissions import IsAdminUser
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.documents.models import Invoice, InvoicePayment, PaymentMethod, PurchaseOrder, PurchaseOrderPayment, Quote
from apps.orders.models import Order
from apps.stores.models import PointOfSale, Stock

from .models import DailyClosing

METHOD_LABELS = dict(PaymentMethod.choices)
ONLINE_LABEL = "En ligne (non affecte)"


def _f(value):
    return float(value or 0)


def _period(request):
    """start/end (YYYY-MM-DD) - par defaut le mois en cours - et point de vente optionnel."""
    today = timezone.localdate()
    p = request.query_params
    try:
        start = date_cls.fromisoformat(p["start"]) if p.get("start") else today.replace(day=1)
        end = date_cls.fromisoformat(p["end"]) if p.get("end") else today
    except ValueError:
        start, end = today.replace(day=1), today
    return start, end, (p.get("point_of_sale") or None)


def _paid_orders(start, end, store):
    qs = Order.objects.filter(status=Order.Status.PAID, paid_at__date__gte=start, paid_at__date__lte=end)
    return qs.filter(point_of_sale_id=store) if store else qs


def _invoices(start, end, store):
    qs = Invoice.objects.filter(status=Invoice.Status.ISSUED, date__gte=start, date__lte=end)
    qs = qs.select_related("customer", "point_of_sale").prefetch_related("items", "payments")
    return qs.filter(point_of_sale_id=store) if store else qs


def _group(qs, key, label=None):
    rows = qs.values(key).annotate(revenue=Sum("total_amount"), orders_count=Count("id")).order_by("-revenue")
    return [
        {"key": r[key], "label": label(r[key]) if label else r[key], "revenue": _f(r["revenue"]), "orders_count": r["orders_count"]}
        for r in rows
    ]


class OverviewReportView(APIView):
    """GET /api/reports/overview/?start=&end=&point_of_sale= - synthese complete d'une periode."""

    permission_classes = [IsAdminUser]

    def get(self, request):
        start, end, store = _period(request)
        orders = _paid_orders(start, end, store)
        agg = orders.aggregate(revenue=Sum("total_amount"), n=Count("id"))
        revenue, n = _f(agg["revenue"]), agg["n"] or 0

        by_day = (
            orders.annotate(day=TruncDate("paid_at")).values("day").annotate(revenue=Sum("total_amount"), n=Count("id")).order_by("day")
        )
        by_store = [
            {**r, "label": r["label"] or ONLINE_LABEL}
            for r in [
                {"label": x["point_of_sale__name"], "revenue": _f(x["revenue"]), "orders_count": x["n"]}
                for x in orders.values("point_of_sale__name").annotate(revenue=Sum("total_amount"), n=Count("id")).order_by("-revenue")
            ]
        ]
        by_cashier = [
            {"label": x["cashier__username"], "revenue": _f(x["revenue"]), "orders_count": x["n"]}
            for x in orders.filter(channel=Order.Channel.POS)
            .values("cashier__username")
            .annotate(revenue=Sum("total_amount"), n=Count("id"))
            .order_by("-revenue")
        ]

        invoices = list(_invoices(start, end, store))
        inv_total = sum((i.total for i in invoices), Decimal("0"))
        inv_paid = sum((i.paid_amount for i in invoices), Decimal("0"))
        pay_qs = InvoicePayment.objects.filter(date__gte=start, date__lte=end, invoice__status=Invoice.Status.ISSUED)
        if store:
            pay_qs = pay_qs.filter(invoice__point_of_sale_id=store)
        inv_payments = defaultdict(Decimal)
        for p in pay_qs:
            inv_payments[p.method] += p.amount

        quotes = list(Quote.objects.filter(date__gte=start, date__lte=end).select_related("customer").prefetch_related("items"))
        if store:
            quotes = [q for q in quotes if str(q.point_of_sale_id) == str(store)]
        q_by_status = defaultdict(int)
        for q in quotes:
            q_by_status[q.status] += 1
        accepted_value = sum((q.total for q in quotes if q.status == Quote.Status.ACCEPTED), Decimal("0"))

        pos = list(
            PurchaseOrder.objects.filter(date__gte=start, date__lte=end)
            .exclude(status=PurchaseOrder.Status.CANCELLED)
            .prefetch_related("items", "payments")
        )
        if store:
            pos = [p for p in pos if str(p.point_of_sale_id) == str(store)]
        po_pay = defaultdict(Decimal)
        for p in PurchaseOrderPayment.objects.filter(date__gte=start, date__lte=end):
            po_pay[p.method] += p.amount

        cash_in = defaultdict(Decimal)
        for r in orders.values("payment_method").annotate(t=Sum("total_amount")):
            cash_in[r["payment_method"]] += r["t"] or Decimal("0")
        for m, v in inv_payments.items():
            cash_in[m] += v

        closings = list(DailyClosing.objects.filter(date__gte=start, date__lte=end))
        if store:
            closings = [c for c in closings if str(c.point_of_sale_id) == str(store)]

        return Response(
            {
                "period": {"start": start, "end": end, "point_of_sale": store},
                "sales": {
                    "revenue": revenue,
                    "orders_count": n,
                    "average_ticket": round(revenue / n, 2) if n else 0,
                    "by_payment_method": _group(orders, "payment_method", lambda k: METHOD_LABELS.get(k, k)),
                    "by_channel": _group(orders, "channel", lambda k: "Caisse (magasin)" if k == "pos" else "En ligne"),
                    "by_store": by_store,
                    "by_cashier": by_cashier,
                    "by_day": [{"day": r["day"], "revenue": _f(r["revenue"]), "orders_count": r["n"]} for r in by_day],
                },
                "invoices": {
                    "count": len(invoices),
                    "invoiced_total": _f(inv_total),
                    "paid_on_these_invoices": _f(inv_paid),
                    "outstanding": _f(inv_total - inv_paid),
                    "collected_in_period_by_method": [
                        {"key": m, "label": METHOD_LABELS.get(m, m), "amount": _f(v)} for m, v in inv_payments.items()
                    ],
                },
                "quotes": {
                    "count": len(quotes),
                    "total": _f(sum((q.total for q in quotes), Decimal("0"))),
                    "by_status": dict(q_by_status),
                    "accepted_value": _f(accepted_value),
                    "conversion_rate": round(100 * q_by_status[Quote.Status.ACCEPTED] / len(quotes), 1) if quotes else 0,
                },
                "purchases": {
                    "count": len(pos),
                    "ordered_total": _f(sum((p.total for p in pos), Decimal("0"))),
                    "paid_in_period_by_method": [
                        {"key": m, "label": METHOD_LABELS.get(m, m), "amount": _f(v)} for m, v in po_pay.items()
                    ],
                },
                "cash_in_by_method": [
                    {"key": m, "label": METHOD_LABELS.get(m, m), "amount": _f(v)}
                    for m, v in sorted(cash_in.items(), key=lambda kv: -kv[1])
                ],
                "cash_discrepancies": {
                    "closings_count": len(closings),
                    "with_discrepancy": sum(1 for c in closings if c.discrepancy_total != 0),
                    "shortage": _f(sum((c.discrepancy_total for c in closings if c.discrepancy_total < 0), Decimal("0"))),
                    "surplus": _f(sum((c.discrepancy_total for c in closings if c.discrepancy_total > 0), Decimal("0"))),
                    "net": _f(sum((c.discrepancy_total for c in closings), Decimal("0"))),
                },
            }
        )


class InventoryReportView(APIView):
    """GET /api/reports/inventory/?point_of_sale=&threshold=5 - stock par magasin, valorisation, ruptures."""

    permission_classes = [IsAdminUser]

    def get(self, request):
        store = request.query_params.get("point_of_sale")
        threshold = int(request.query_params.get("threshold", 5))
        stocks = Stock.objects.select_related("product", "point_of_sale", "product__category").filter(
            product__is_active=True, point_of_sale__is_active=True
        )
        if store:
            stocks = stocks.filter(point_of_sale_id=store)

        per_store = defaultdict(lambda: {"units": 0, "value": Decimal("0"), "references": 0})
        rows, low = [], []
        for s in stocks:
            value = s.quantity * s.product.price
            agg = per_store[s.point_of_sale.name]
            agg["units"] += s.quantity
            agg["value"] += value
            agg["references"] += 1
            row = {
                "store": s.point_of_sale.name,
                "sku": s.product.sku,
                "product": s.product.name,
                "category": s.product.category.name if s.product.category else None,
                "quantity": s.quantity,
                "unit_price": _f(s.product.price),
                "value": _f(value),
            }
            rows.append(row)
            if s.quantity <= threshold:
                low.append(row)
        rows.sort(key=lambda r: (r["store"], r["product"]))
        low.sort(key=lambda r: (r["quantity"], r["product"]))
        return Response(
            {
                "threshold": threshold,
                "note": "Valorisation au prix de vente actuel (le prix d'achat n'est pas enregistre sur les produits).",
                "totals": {
                    "units": sum(v["units"] for v in per_store.values()),
                    "value": _f(sum((v["value"] for v in per_store.values()), Decimal("0"))),
                },
                "by_store": [
                    {"store": k, "units": v["units"], "value": _f(v["value"]), "references": v["references"]}
                    for k, v in sorted(per_store.items())
                ],
                "low_stock": low,
                "rows": rows,
            }
        )


def _aging(days_overdue):
    if days_overdue <= 0:
        return "Non echu"
    if days_overdue <= 30:
        return "1-30 jours"
    if days_overdue <= 60:
        return "31-60 jours"
    if days_overdue <= 90:
        return "61-90 jours"
    return "Plus de 90 jours"


def _open_documents(qs, party):
    today = timezone.localdate()
    rows, buckets = [], defaultdict(Decimal)
    for d in qs:
        balance = d.balance
        if balance <= 0:
            continue
        due = getattr(d, "due_date", None) or getattr(d, "expected_date", None)
        days = (today - due).days if due else 0
        bucket = _aging(days)
        buckets[bucket] += balance
        rows.append(
            {
                "id": d.id,
                "number": d.number,
                "party": getattr(d, party).name,
                "date": d.date,
                "due_date": due,
                "total": _f(d.total),
                "paid": _f(d.paid_amount),
                "balance": _f(balance),
                "days_overdue": max(days, 0),
                "bucket": bucket,
            }
        )
    rows.sort(key=lambda r: -r["days_overdue"])
    order = ["Non echu", "1-30 jours", "31-60 jours", "61-90 jours", "Plus de 90 jours"]
    return {
        "total_outstanding": _f(sum(buckets.values(), Decimal("0"))),
        "aging": [{"bucket": b, "amount": _f(buckets.get(b))} for b in order],
        "rows": rows,
    }


class ReceivablesView(APIView):
    """GET /api/reports/receivables/ - factures clients impayees (creances) par anciennete."""

    permission_classes = [IsAdminUser]

    def get(self, request):
        qs = Invoice.objects.filter(status=Invoice.Status.ISSUED).select_related("customer").prefetch_related("items", "payments")
        return Response(_open_documents(qs, "customer"))


class PayablesView(APIView):
    """GET /api/reports/payables/ - bons de commande fournisseurs non soldes (dettes)."""

    permission_classes = [IsAdminUser]

    def get(self, request):
        qs = (
            PurchaseOrder.objects.exclude(status__in=[PurchaseOrder.Status.CANCELLED, PurchaseOrder.Status.DRAFT])
            .select_related("supplier")
            .prefetch_related("items", "payments")
        )
        return Response(_open_documents(qs, "supplier"))


class CashDiscrepancyReportView(APIView):
    """GET /api/reports/cash-discrepancies/?start=&end=&point_of_sale= - fermetures de caisse et ecarts."""

    permission_classes = [IsAdminUser]

    def get(self, request):
        start, end, store = _period(request)
        qs = DailyClosing.objects.filter(date__gte=start, date__lte=end).select_related("point_of_sale", "cashier", "closed_by")
        if store:
            qs = qs.filter(point_of_sale_id=store)
        rows = [
            {
                "id": c.id,
                "date": c.date,
                "store": c.point_of_sale.name if c.point_of_sale else ONLINE_LABEL,
                "cashier": c.cashier.username if c.cashier else None,
                "expected": _f(c.expected_total),
                "declared": _f(c.declared_total),
                "discrepancy": _f(c.discrepancy_total),
                "initial_discrepancy": _f(c.initial_discrepancy_total) if c.cashier else None,
                "corrections": c.revision_count,
                "notes": c.notes,
            }
            for c in qs
        ]
        return Response(
            {
                "rows": rows,
                "shortage": sum(r["discrepancy"] for r in rows if r["discrepancy"] < 0),
                "surplus": sum(r["discrepancy"] for r in rows if r["discrepancy"] > 0),
                "net": sum(r["discrepancy"] for r in rows),
            }
        )


class ExportView(APIView):
    """GET /api/reports/export/?start=&end=&point_of_sale= - classeur Excel multi-feuilles."""

    permission_classes = [IsAdminUser]

    def get(self, request):
        start, end, store = _period(request)
        wb = openpyxl.Workbook()

        def sheet(title, headers, rows):
            ws = wb.create_sheet(title)
            ws.append(headers)
            for r in rows:
                ws.append(r)
            for i, h in enumerate(headers, start=1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(h) + 4)

        wb.remove(wb.active)
        orders = _paid_orders(start, end, store).select_related("point_of_sale", "cashier").order_by("paid_at")
        sheet(
            "Ventes",
            ["Reference", "Date", "Canal", "Point de vente", "Caissier", "Client", "Mode de paiement", "Total (FCFA)"],
            [
                [
                    o.reference[:8].upper(),
                    timezone.localtime(o.paid_at).strftime("%Y-%m-%d %H:%M") if o.paid_at else "",
                    o.get_channel_display(),
                    o.point_of_sale.name if o.point_of_sale else ONLINE_LABEL,
                    o.cashier.username if o.cashier else "",
                    o.customer_name,
                    METHOD_LABELS.get(o.payment_method, o.payment_method),
                    _f(o.total_amount),
                ]
                for o in orders
            ],
        )
        invoices = list(_invoices(start, end, store))
        sheet(
            "Factures",
            ["Numero", "Date", "Echeance", "Client", "Point de vente", "Statut paiement", "Total", "Paye", "Solde"],
            [
                [i.number, i.date, i.due_date, i.customer.name, i.point_of_sale.name if i.point_of_sale else "", i.payment_status, _f(i.total), _f(i.paid_amount), _f(i.balance)]
                for i in invoices
            ],
        )
        sheet(
            "Paiements factures",
            ["Facture", "Date", "Client", "Mode de paiement", "Montant", "Reference"],
            [
                [p.invoice.number, p.date, p.invoice.customer.name, METHOD_LABELS.get(p.method, p.method), _f(p.amount), p.reference]
                for p in InvoicePayment.objects.filter(date__gte=start, date__lte=end).select_related("invoice", "invoice__customer")
            ],
        )
        quotes = Quote.objects.filter(date__gte=start, date__lte=end).select_related("customer", "point_of_sale").prefetch_related("items")
        sheet(
            "Devis",
            ["Numero", "Date", "Valable jusqu'au", "Client", "Statut", "Total"],
            [[q.number, q.date, q.valid_until, q.customer.name, q.get_status_display(), _f(q.total)] for q in quotes],
        )
        pos = PurchaseOrder.objects.filter(date__gte=start, date__lte=end).select_related("supplier").prefetch_related("items", "payments")
        sheet(
            "Bons de commande",
            ["Numero", "Date", "Livraison prevue", "Fournisseur", "Statut", "Mode de paiement", "Total", "Paye", "Solde"],
            [
                [p.number, p.date, p.expected_date, p.supplier.name, p.get_status_display(), p.get_payment_method_display(), _f(p.total), _f(p.paid_amount), _f(p.balance)]
                for p in pos
            ],
        )
        closings = DailyClosing.objects.filter(date__gte=start, date__lte=end).select_related("point_of_sale", "cashier")
        sheet(
            "Ecarts de caisse",
            ["Date", "Point de vente", "Caissier", "Attendu", "Declare", "Ecart", "Ecart initial", "Corrections", "Notes"],
            [
                [c.date, c.point_of_sale.name if c.point_of_sale else ONLINE_LABEL, c.cashier.username if c.cashier else "", _f(c.expected_total), _f(c.declared_total), _f(c.discrepancy_total), _f(c.initial_discrepancy_total), c.revision_count, c.notes]
                for c in closings
            ],
        )
        sheet(
            "Stock",
            ["Point de vente", "Reference", "Produit", "Quantite", "Prix de vente", "Valeur"],
            [
                [s.point_of_sale.name, s.product.sku, s.product.name, s.quantity, _f(s.product.price), _f(s.quantity * s.product.price)]
                for s in Stock.objects.select_related("product", "point_of_sale").order_by("point_of_sale__name", "product__name")
            ],
        )

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="rapports_labelleteranga_{start}_{end}.xlsx"'
        return response
